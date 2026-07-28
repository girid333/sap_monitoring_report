import os, copy
from datetime import datetime
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelHandler:
    def get_variable_fields(self, steps: List[Dict]) -> List[Dict]:
        variable_fields = []
        seen = set()
        for i, step in enumerate(steps):
            if step.get('type') in ('fill_by_label', 'fill_by_id', 'fill_by_position'):
                field_name = step.get('label') or step.get('element_id') or f'Field_{i+1}'
                if field_name not in seen:
                    seen.add(field_name)
                    variable_fields.append({
                        'step_index': i,
                        'field_name': field_name,
                        'default_value': step.get('value', '')
                    })
        return variable_fields

    def generate_template(self, job: Dict) -> str:
        variable_fields = self.get_variable_fields(job.get('steps', []))
        
        wb = openpyxl.Workbook()
        ws_data = wb.active
        ws_data.title = 'Data Input'
        ws_inst = wb.create_sheet('Instructions')
        
        ws_inst['A1'] = 'How to use this template:'
        ws_inst['A1'].font = Font(bold=True, size=14, color='1E3A5F')
        
        ws_inst['A3'] = "1. Each row in the 'Data Input' sheet represents one run of the recording."
        ws_inst['A4'] = "2. Fill in the 'Run_Name' to identify each run."
        ws_inst['A5'] = "3. Fill in the specific values for each variable field for that run."
        ws_inst['A6'] = "4. Do not modify the header row (Row 1)."
        ws_inst['A7'] = "5. Save the file and upload it to run the batch."
        ws_inst['A8'] = "6. You can add as many rows as you need for batch processing."
            
        ws_inst['A10'] = f"Job: {job.get('job_name', '')}"
        ws_inst['A11'] = f"Variable Fields: {len(variable_fields)}"
        ws_inst['A12'] = f"Total Steps in Recording: {len(job.get('steps', []))}"
        
        headers = ['Run_Name'] + [f['field_name'] for f in variable_fields]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws_data.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1E3A5F')
            cell.alignment = Alignment(horizontal='center')
            
        example_data = ['Example_Run_1'] + [f['default_value'] for f in variable_fields]
        for col_idx, val in enumerate(example_data, start=1):
            cell = ws_data.cell(row=2, column=col_idx, value=val)
            cell.fill = PatternFill('solid', fgColor='EFF6FF')
            
        ws_data.freeze_panes = 'A2'
        
        for col_idx, header in enumerate(headers, start=1):
            col_letter = get_column_letter(col_idx)
            ws_data.column_dimensions[col_letter].width = max(18, len(str(header)) + 4)
            
        if not variable_fields:
            ws_data['A3'] = 'No variable fields found in this recording. All steps use fixed values.'
            
        safe_job_name = job.get('job_name', '').replace(' ', '_').replace('/', '_')
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        save_dir = os.path.join(os.path.dirname(__file__), 'custom_recordings')
        os.makedirs(save_dir, exist_ok=True)
        
        save_path = os.path.join(save_dir, f"template_{safe_job_name}_{timestamp}.xlsx")
        wb.save(save_path)
        return os.path.abspath(save_path)

    def read_batch_data(self, excel_path: str) -> List[Dict]:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        if 'Data Input' in wb.sheetnames:
            ws = wb['Data Input']
        else:
            ws = wb.active
            
        headers = [str(cell.value) for cell in ws[1] if cell.value is not None]
        if not headers:
            return []
            
        batch_data = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(cell is None or str(cell).strip() == '' for cell in row):
                continue
            row_dict = {}
            for header, cell_value in zip(headers, row):
                row_dict[header] = str(cell_value) if cell_value is not None else ''
            batch_data.append(row_dict)
            
        return batch_data

    def substitute_values(self, steps: List[Dict], row_data: Dict) -> List[Dict]:
        modified_steps = copy.deepcopy(steps)
        for step in modified_steps:
            if step.get('type') in ('fill_by_label', 'fill_by_id', 'fill_by_position'):
                field_name = step.get('label') or step.get('element_id') or ''
                if field_name in row_data:
                    step['value'] = row_data[field_name]
        return modified_steps
