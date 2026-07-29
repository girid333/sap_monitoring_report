import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from typing import List, Dict, Any

class CustomReportGenerator:
    def __init__(self, output_dir):
        self.output_dir = output_dir
        if not os.path.exists(self.output_dir):
            os.makedirs(self.output_dir)

    def _set_cell_background(self, cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = tcPr.first_child_found_in("w:tcBorders")
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    def _add_horizontal_rule(self, doc):
        p = doc.add_paragraph()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), 'auto')
        pBdr.append(bottom)
        p._p.get_or_add_pPr().append(pBdr)

    def format_step_description(self, step) -> str:
        step_type = step.get('type', 'Unknown step')
        if step_type == 'navigate_tcode':
            return f"Navigate to T-Code: {step.get('tcode')}"
        elif step_type == 'press_fkey':
            return f"Press Key: {step.get('key')}"
        elif step_type == 'fill_by_label':
            return f"Fill Field '{step.get('label')}' with '{step.get('value')}'"
        elif step_type == 'fill_by_id':
            return f"Fill Element ID '{step.get('element_id')}' with '{step.get('value')}'"
        elif step_type == 'fill_by_position':
            return f"Fill field at position {step.get('position')} with '{step.get('value')}'"
        elif step_type == 'select_dropdown':
            return f"Select '{step.get('option')}' from dropdown '{step.get('label_or_id', '')}'"
        elif step_type == 'clear_field':
            return f"Clear field '{step.get('label_or_id', '')}'"
        elif step_type == 'click_button':
            return f"Click button: '{step.get('button_text')}'"
        elif step_type == 'click_menu_path':
            return "Menu: " + " → ".join(step.get('path', []))
        elif step_type == 'click_tab':
            return f"Click tab: '{step.get('tab_text')}'"
        elif step_type == 'click_table_row':
            return f"Click table row containing: '{step.get('row_text')}'"
        elif step_type == 'double_click':
            return f"Double-click: '{step.get('target', '')}'"
        elif step_type == 'expand_tree_node':
            return f"Expand tree node: '{step.get('node_text')}'"
        elif step_type == 'click_tree_node':
            return f"Click tree node: '{step.get('node_text')}'"
        elif step_type == 'confirm_dialog':
            return "Confirm dialog (click OK/Yes)"
        elif step_type == 'dismiss_dialog':
            return "Dismiss dialog (click Cancel/No)"
        elif step_type == 'handle_f4_help':
            return f"F4 value help for '{step.get('field_label')}', search: '{step.get('search_value')}'"
        elif step_type == 'scroll_table_down':
            return f"Scroll table down {step.get('times', 1)} time(s)"
        elif step_type == 'filter_column':
            return f"Filter column '{step.get('column_name')}' by '{step.get('filter_value')}'"
        elif step_type == 'wait_seconds':
            return f"Wait {step.get('seconds', 1)} second(s)"
        elif step_type == 'wait_for_element':
            return f"Wait for element: '{step.get('text')}'"
        elif step_type == 'scroll_page_down':
            return "Scroll page down"
        elif step_type == 'screenshot':
            return f"📸 Screenshot: '{step.get('caption', '')}'"
        elif step_type == 'screenshot_full_page':
            return f"📸 Full-page screenshot: '{step.get('caption', '')}'"
        elif step_type == 'for_each_table_row':
            return f"Loop over table rows (max {step.get('max_rows', 10)} rows)"
        return step_type

    def generate(self, job_name, job_description, system_name, sid, steps, screenshots, output_filename=None) -> str:
        if not output_filename:
            safe_job = job_name.replace(' ', '_')
            ts = datetime.now().strftime('%Y%m%d_%H%M%S')
            output_filename = f"custom_recording_{safe_job}_{ts}.docx"
            
        doc_path = os.path.join(self.output_dir, output_filename)
        doc = Document()
        
        # 1. Cover Page Section
        p1 = doc.add_paragraph()
        run1 = p1.add_run("SAP Custom T-Code Recording")
        run1.font.size = Pt(32)
        run1.font.bold = True
        run1.font.color.rgb = RGBColor(0, 70, 127)
        
        p2 = doc.add_paragraph()
        run2 = p2.add_run(f"Job: {job_name}")
        run2.font.size = Pt(18)
        run2.font.bold = True
        
        p3 = doc.add_paragraph()
        run3 = p3.add_run(f"System: {system_name} ({sid})")
        run3.font.size = Pt(14)
        
        p4 = doc.add_paragraph()
        run4 = p4.add_run(f"Generated: {datetime.now().strftime('%d %B %Y %H:%M')}")
        run4.font.size = Pt(12)
        
        if job_description:
            p5 = doc.add_paragraph()
            run5 = p5.add_run(job_description)
            run5.font.size = Pt(11)
            run5.font.italic = True
            
        self._add_horizontal_rule(doc)
        doc.add_page_break()
        
        # 2. Job Summary Section
        doc.add_heading("Recording Summary", level=1)
        
        summary_data = [
            ("Job Name", job_name),
            ("System", system_name),
            ("SID", sid),
            ("Total Steps", str(len(steps))),
            ("Screenshots Captured", str(len(screenshots))),
            ("Generated On", datetime.now().strftime('%d %B %Y %H:%M'))
        ]
        
        table = doc.add_table(rows=len(summary_data), cols=2)
        for i, (key, value) in enumerate(summary_data):
            row_cells = table.rows[i].cells
            row_cells[0].text = key
            row_cells[1].text = value
            
        doc.add_page_break()
        
        # 3. Step-by-Step Recording Section
        doc.add_heading("Step-by-Step Recording", level=1)
        
        for i, step in enumerate(steps):
            step_text = f"Step {i+1}: {self.format_step_description(step)}"
            p_step = doc.add_paragraph()
            run_step = p_step.add_run(step_text)
            run_step.font.bold = True
            
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), 'EEEEEE')
            p_step._p.get_or_add_pPr().append(shd)
            
            step_screenshot = None
            for ss in screenshots:
                if ss.get('step_index') == i or (step.get('caption') and ss.get('caption') == step.get('caption')):
                    step_screenshot = ss
                    break
                    
            img_path = step_screenshot.get('path', step_screenshot.get('filepath')) if step_screenshot else None
            if img_path and os.path.exists(img_path):
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_run = img_p.add_run()
                img_run.add_picture(img_path, width=Inches(6.0))
                
                cap_p = doc.add_paragraph()
                cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_run = cap_p.add_run(step_screenshot.get('caption', f"Screenshot for Step {i+1}"))
                cap_run.font.italic = True
                cap_run.font.color.rgb = RGBColor(128, 128, 128)
                
            self._add_horizontal_rule(doc)
            
        # 4. Screenshot Gallery Section
        if screenshots:
            doc.add_page_break()
            doc.add_heading("Screenshot Gallery", level=1)
            
            gallery_table = doc.add_table(rows=0, cols=2)
            gallery_table.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            for i in range(0, len(screenshots), 2):
                row = gallery_table.add_row()
                
                # Column 1
                cell1 = row.cells[0]
                ss1 = screenshots[i]
                p_c1 = cell1.add_paragraph()
                p_c1.add_run(f"Screenshot {i+1}").font.bold = True
                
                img_path1 = ss1.get('path', ss1.get('filepath'))
                if img_path1 and os.path.exists(img_path1):
                    p_img1 = cell1.add_paragraph()
                    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img1.add_run().add_picture(img_path1, width=Inches(5.5))
                
                p_cap1 = cell1.add_paragraph()
                p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_cap1 = p_cap1.add_run(ss1.get('caption', ''))
                run_cap1.font.italic = True
                
                # Column 2
                if i + 1 < len(screenshots):
                    cell2 = row.cells[1]
                    ss2 = screenshots[i+1]
                    p_c2 = cell2.add_paragraph()
                    p_c2.add_run(f"Screenshot {i+2}").font.bold = True
                    
                    img_path2 = ss2.get('path', ss2.get('filepath'))
                    if img_path2 and os.path.exists(img_path2):
                        p_img2 = cell2.add_paragraph()
                        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img2.add_run().add_picture(img_path2, width=Inches(5.5))
                        
                    p_cap2 = cell2.add_paragraph()
                    p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run_cap2 = p_cap2.add_run(ss2.get('caption', ''))
                    run_cap2.font.italic = True

        doc.save(doc_path)
        return doc_path

    def generate_batch(self, job_name: str, system_name: str, sid: str,
                       all_results: List[Dict]) -> str:
        """
        Generate a combined Word report for all batch runs.
        all_results: list of dicts with keys: run_name, row_data, success, screenshots
        screenshots: list of dicts with keys: path, caption
        """
        doc = Document()

        # Cover
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run('SAP Custom T-Code — Batch Recording Report')
        run.bold = True
        run.font.size = Pt(26)
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

        doc.add_paragraph(f'Job: {job_name}').runs[0].font.size = Pt(14)
        doc.add_paragraph(f'System: {system_name} ({sid})').runs[0].font.size = Pt(12)
        doc.add_paragraph(f'Generated: {datetime.now().strftime("%d %B %Y %H:%M")}').runs[0].font.size = Pt(11)
        doc.add_paragraph(f'Total Runs: {len(all_results)}').runs[0].font.size = Pt(11)
        self._add_horizontal_rule(doc)
        doc.add_page_break()

        # One section per run
        for idx, result in enumerate(all_results):
            run_name = result.get('run_name', f'Run {idx+1}')
            row_data = result.get('row_data', {})
            success = result.get('success', False)
            screenshots = result.get('screenshots', [])

            heading = doc.add_heading(f'Run {idx+1}: {run_name}', level=1)
            heading.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

            # Status badge
            status_p = doc.add_paragraph()
            status_run = status_p.add_run('✓ SUCCESS' if success else '✗ FAILED')
            status_run.bold = True
            status_run.font.color.rgb = RGBColor(0x16, 0x65, 0x34) if success else RGBColor(0x99, 0x1B, 0x1B)

            # Input data table
            if row_data:
                doc.add_paragraph('Input Data:', style='Heading 3')
                data_items = {k: v for k, v in row_data.items() if k != 'Run_Name'}
                if data_items:
                    tbl = doc.add_table(rows=1, cols=2)
                    tbl.style = 'Table Grid'
                    hdr = tbl.rows[0].cells
                    hdr[0].text = 'Field'
                    hdr[1].text = 'Value'
                    for cell in hdr:
                        cell.paragraphs[0].runs[0].bold = True
                        self._set_cell_background(cell, 'DBEAFE')
                    for field_name, field_val in data_items.items():
                        row_cells = tbl.add_row().cells
                        row_cells[0].text = str(field_name)
                        row_cells[1].text = str(field_val)

            # Screenshots
            if screenshots:
                doc.add_paragraph('Screenshots:', style='Heading 3')
                for ss in screenshots:
                    ss_path = ss.get('path', '')
                    ss_caption = ss.get('caption', '')
                    if ss_path and os.path.exists(ss_path):
                        p = doc.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.add_run().add_picture(ss_path, width=Inches(6.0))
                        cap_p = doc.add_paragraph()
                        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cap_run = cap_p.add_run(ss_caption)
                        cap_run.italic = True
                        cap_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            else:
                doc.add_paragraph('No screenshots captured for this run.')

            if idx < len(all_results) - 1:
                doc.add_page_break()

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        safe_name = job_name.replace(' ', '_').replace('/', '_')
        doc_path = os.path.join(self.output_dir,
                                f'batch_recording_{safe_name}_{timestamp}.docx')
        doc.save(doc_path)
        return doc_path
